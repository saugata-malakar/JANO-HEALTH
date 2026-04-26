"""Resume file → plain text extraction.

Supports PDF (via pypdf) and DOCX (via python-docx). For PDFs that are
purely scanned images we fall back to noting that OCR is needed; the
implementation hook is documented below.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class ResumeExtractionError(Exception):
    pass


def extract_text(filename: str, content: bytes) -> str:
    """Dispatch to the right extractor based on file extension."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(content)
    if ext == ".docx":
        return _extract_docx(content)
    if ext in {".txt", ".md"}:
        return content.decode("utf-8", errors="replace")
    raise ResumeExtractionError(f"Unsupported file type: {ext}")


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception as e:  # noqa: BLE001
                logger.warning("PDF page extraction failed: %s", e)
        text = "\n\n".join(p for p in pages if p.strip())
    except Exception as e:
        raise ResumeExtractionError(f"Failed to read PDF: {e}") from e

    if not text.strip():
        # Likely a scanned PDF — would need OCR (Tesseract/PaddleOCR) in prod.
        raise ResumeExtractionError(
            "PDF contains no extractable text. It may be a scanned image — "
            "OCR support is not enabled in this build."
        )
    return _normalize(text)


def _extract_docx(content: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(content))
    except Exception as e:
        raise ResumeExtractionError(f"Failed to read DOCX: {e}") from e

    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Tables often hold contact info / skills grids
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _normalize("\n".join(parts))


def _normalize(text: str) -> str:
    """Collapse runs of whitespace and trim each line."""
    lines = [line.strip() for line in text.splitlines()]
    # Drop empty runs longer than 1
    cleaned: list[str] = []
    blank = False
    for line in lines:
        if not line:
            if not blank:
                cleaned.append("")
            blank = True
        else:
            cleaned.append(line)
            blank = False
    return "\n".join(cleaned).strip()
